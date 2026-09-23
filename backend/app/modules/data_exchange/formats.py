"""Bounded file parsing and reversible, formula-safe serialization with Pandas."""

import csv
import io
import json
import math
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from xml.etree.ElementTree import ParseError
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile, ZipInfo

import pandas as pd
from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from sqlalchemy import JSON, Boolean, Date, DateTime, Enum, Integer, Numeric, String, Time

from app.modules.data_exchange.registry import FORMAT_VERSION, TABLES, columns_for

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNPACKED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 100_000
MAX_TOTAL_ROWS = 250_000
MAX_CELL_CHARS = 1_000_000
NULL = "\\N"
EMPTY = "\\E"


class ExchangeError(ValueError):
    def __init__(
        self,
        message: str,
        table: str | None = None,
        row: int | None = None,
        column: str | None = None,
    ):
        super().__init__(message)
        self.detail = {"message": message, "table": table, "row": row, "column": column}


class ParsedRow(dict):
    """Keep the source row for database errors without changing package fingerprints."""

    def __init__(self, row_number: int):
        super().__init__()
        self.row_number = row_number


def json_default(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(type(value).__name__)


def encode_cell(value) -> str:
    if value is None:
        return NULL
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    elif isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    elif isinstance(value, bool):
        value = "true" if value else "false"
    else:
        value = str(value)
    if value == "":
        return EMPTY
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        return "\\T" + json.dumps(value, ensure_ascii=True)
    # Escape strings so opening CSV in a spreadsheet cannot execute a formula.
    # Prefix backslashes too, preserving the distinction between NULL and literal \N.
    if value.startswith(("'", "\\")) or value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    if value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _decode_cell(value):
    if value is None or value == NULL or pd.isna(value):
        return None
    if value == EMPTY:
        return ""
    if isinstance(value, str):
        if value.startswith('\\T"'):
            return json.loads(value[2:])
        if value.startswith("'"):
            rest = value[1:]
            if rest.startswith(("'", "\\", "\t", "\r", "\n")) or rest.lstrip().startswith(
                ("=", "+", "-", "@")
            ):
                return rest
    return value


def decode_cell(value):
    return _decode_cell(value)


def _convert_type(value, column):
    kind = column.type
    if isinstance(kind, Boolean):
        if value in (True, "true", "1", 1):
            return True
        if value in (False, "false", "0", 0):
            return False
        raise ValueError("Ожидается true или false")
    if isinstance(kind, Integer):
        if isinstance(value, bool):
            raise ValueError("Ожидается целое число")
        numeric = Decimal(str(value))
        if not numeric.is_finite() or numeric != numeric.to_integral_value():
            raise ValueError("Ожидается целое число")
        result = int(numeric)
        if not -2_147_483_648 <= result <= 2_147_483_647:
            raise ValueError("Целое число выходит за диапазон int32")
        if (column.primary_key or column.foreign_keys) and result <= 0:
            raise ValueError("Идентификатор должен быть положительным")
        return result
    if isinstance(kind, Numeric):
        result = Decimal(str(value))
        if not result.is_finite():
            raise ValueError("Ожидается конечное число")
        if kind.precision is not None and kind.scale is not None:
            if abs(result) >= Decimal(10) ** (kind.precision - kind.scale):
                raise ValueError("Число выходит за точность столбца")
            if result != result.quantize(Decimal(10) ** -kind.scale):
                raise ValueError(f"Допускается не более {kind.scale} знаков после запятой")
        return result
    if isinstance(kind, DateTime):
        result = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
        if result.tzinfo is None or result.utcoffset() is None:
            raise ValueError("Время должно содержать часовой пояс, например +03:00")
        return result
    if isinstance(kind, Date):
        if isinstance(value, datetime):
            if value.time() != time():
                raise ValueError("В дате маршрута не должно быть времени")
            return value.date()
        return value if isinstance(value, date) else date.fromisoformat(str(value))
    if isinstance(kind, Time):
        result = value if isinstance(value, time) else time.fromisoformat(str(value))
        if result.tzinfo is not None:
            raise ValueError("Время смены задаётся без часового пояса")
        return result
    if isinstance(kind, JSON):
        result = json.loads(value) if isinstance(value, str) else value
        # Reject NaN/Infinity at any nesting level.
        json.dumps(result, allow_nan=False)
        if not isinstance(result, dict):
            raise ValueError("Ожидается JSON-объект")
        return result
    result = str(value)
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Недопустимое число")
    if isinstance(kind, Enum) and result not in kind.enums:
        raise ValueError(f"Допустимые значения: {', '.join(kind.enums)}")
    if isinstance(kind, String) and kind.length and len(result) > kind.length:
        raise ValueError(f"Максимальная длина: {kind.length}")
    return result


def convert(value, column):
    value = decode_cell(value)
    if value is None or (value == "" and not isinstance(column.type, String)):
        if not column.nullable:
            raise ValueError("Обязательное поле не может быть пустым")
        return None
    if isinstance(value, str) and ("\x00" in value or len(value) > MAX_CELL_CHARS):
        raise ValueError("Недопустимый символ или слишком длинное значение")
    return _convert_type(value, column)


def normalize_dataframe(
    name: str, df: pd.DataFrame, source_row_numbers: pd.Series
) -> list[ParsedRow]:
    columns = {c.name: c for c in columns_for(name)}
    converted_columns: dict[str, list] = {}

    for key in df.columns:
        col = columns[key]
        series = df[key]

        # 1. Cell length and null byte validation
        def check_cell(v):
            return not (isinstance(v, str) and ("\x00" in v or len(v) > MAX_CELL_CHARS))

        valid_chars = series.map(check_cell)
        if not valid_chars.all():
            bad_idx = (~valid_chars).idxmax()
            bad_row = int(source_row_numbers.loc[bad_idx])
            raise ExchangeError(
                "Недопустимый символ или слишком длинное значение", name, bad_row, key
            )

        # 2. Decode cell representation
        decoded = series.map(_decode_cell)

        # 3. Check nullable constraints
        def _is_empty(v):
            return v is None or (isinstance(v, float) and pd.isna(v)) or (
                v == "" and not isinstance(col.type, String)
            )

        is_empty = decoded.map(_is_empty)
        if not col.nullable and is_empty.any():
            bad_idx = is_empty.idxmax()
            bad_row = int(source_row_numbers.loc[bad_idx])
            raise ExchangeError("Обязательное поле не может быть пустым", name, bad_row, key)

        # 4. Convert and validate column values into pure Python objects
        converted = []
        for idx, val in decoded.items():
            if val is None or (isinstance(val, float) and pd.isna(val)) or (
                val == "" and not isinstance(col.type, String)
            ):
                converted.append(None)
            else:
                try:
                    converted.append(_convert_type(val, col))
                except (
                    ValueError,
                    TypeError,
                    InvalidOperation,
                    OverflowError,
                    RecursionError,
                ) as error:
                    bad_row = int(source_row_numbers.loc[idx])
                    raise ExchangeError(str(error), name, bad_row, key) from error
        converted_columns[key] = converted

    # 5. Fast vectorized primary key uniqueness validation
    pk_cols = [c.name for c in TABLES[name].primary_key]
    if pk_cols:
        pk_df = pd.DataFrame({col: converted_columns[col] for col in pk_cols})
        dup_mask = pk_df.duplicated(keep="first")
        if dup_mask.any():
            first_dup_idx = dup_mask.idxmax()
            first_dup_row = int(source_row_numbers.iloc[first_dup_idx])
            raise ExchangeError("Повторяющийся первичный ключ", name, first_dup_row)

    # 6. Build ParsedRow list preserving row_number and pure Python types
    result = []
    row_nums = source_row_numbers.to_list()
    col_names = list(df.columns)
    num_rows = len(row_nums)
    for i in range(num_rows):
        row = ParsedRow(row_nums[i])
        for col_name in col_names:
            row[col_name] = converted_columns[col_name][i]
        result.append(row)
    return result


def normalize_rows(name: str, rows) -> list[dict]:
    if name not in TABLES:
        raise ExchangeError("Неизвестная таблица", name)
    iterator = iter(rows)
    header = next(iterator, None)
    if not header:
        raise ExchangeError("Нет строки заголовков", name, 1)
    if not all(isinstance(c, str) and c for c in header) or len(header) != len(set(header)):
        raise ExchangeError("Пустые или повторяющиеся заголовки", name, 1)
    columns = {c.name: c for c in columns_for(name)}
    if set(header) - columns.keys():
        raise ExchangeError(
            "Неизвестные столбцы: " + ", ".join(sorted(set(header) - columns.keys())), name, 1
        )
    required = {
        c.name
        for c in columns.values()
        if not c.nullable and c.server_default is None and c.default is None
    }
    required |= {c.name for c in columns.values() if c.primary_key}
    if required - set(header):
        raise ExchangeError("Нет столбцов: " + ", ".join(sorted(required - set(header))), name, 1)

    raw_data = []
    source_row_numbers = []
    header_len = len(header)
    for row_number, row in enumerate(iterator, 2):
        if row_number > MAX_ROWS + 1:
            raise ExchangeError("Превышен лимит строк", name, row_number)
        if all(value is None or value == "" for value in row):
            continue
        if len(row) != header_len:
            raise ExchangeError("Число ячеек не соответствует заголовку", name, row_number)
        raw_data.append(row)
        source_row_numbers.append(row_number)

    if not raw_data:
        return []

    df = pd.DataFrame(raw_data, columns=header, dtype=object)
    row_nums = pd.Series(source_row_numbers, index=df.index)
    return normalize_dataframe(name, df, row_nums)


def read_csv(content: bytes, name: str) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ExchangeError("CSV должен быть в UTF-8", name) from error

    try:
        first = text.splitlines()[0] if text else ""
        delimiter = max((",", ";", "\t"), key=first.count)
        csv.field_size_limit(MAX_CELL_CHARS)
        return normalize_rows(
            name, csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        )
    except csv.Error as error:
        raise ExchangeError("Некорректный CSV: " + str(error), name) from error


def inspect_archive(content: bytes) -> None:
    with ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 500 or sum(e.file_size for e in entries) > MAX_UNPACKED_BYTES:
            raise ExchangeError("Слишком большой распакованный архив")
        if len({e.filename for e in entries}) != len(entries):
            raise ExchangeError("Повторяющиеся имена файлов в архиве")
        if any(e.flag_bits & 1 for e in entries):
            raise ExchangeError("Зашифрованные архивы не поддерживаются")


def parse_file(content: bytes, filename: str, entity: str | None = None) -> dict[str, list[dict]]:
    if not content or len(content) > MAX_FILE_BYTES:
        raise ExchangeError("Пустой файл или размер больше 20 MiB")
    extension = filename.rsplit(".", 1)[-1].lower()
    try:
        if extension == "csv":
            if entity is None:
                raise ExchangeError("Для CSV укажите параметр entity")
            return {entity: read_csv(content, entity)}
        if extension not in ("zip", "xlsx"):
            raise ExchangeError("Поддерживаются .csv, .zip (CSV-пакет) и .xlsx")
        if entity is not None:
            raise ExchangeError("entity применяется только к одиночному CSV")
        inspect_archive(content)
        result = {}
        if extension == "zip":
            with ZipFile(io.BytesIO(content)) as archive:
                manifest = json.loads(archive.read("manifest.json"))
                if manifest not in ({"format_version": "1"}, {"format_version": FORMAT_VERSION}):
                    raise ExchangeError("Неподдерживаемая версия CSV-пакета")
                for entry in archive.infolist():
                    if entry.filename == "manifest.json":
                        continue
                    if not entry.filename.endswith(".csv") or entry.filename[:-4] not in TABLES:
                        raise ExchangeError("Неизвестный файл в CSV-пакете: " + entry.filename)
                    name = entry.filename[:-4]
                    result[name] = read_csv(archive.read(entry), name)
        else:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                if "_meta" not in workbook:
                    raise ExchangeError("Нет версии формата в листе _meta")
                meta = workbook["_meta"]
                if (
                    (meta.max_row or 0) > 1
                    or (meta.max_column or 0) > 2
                    or list(meta.iter_rows(max_row=1, max_col=2, values_only=True))
                    not in ([("format_version", "1")], [("format_version", FORMAT_VERSION)])
                ):
                    raise ExchangeError("Нет версии формата в листе _meta")
                for sheet in workbook:
                    if sheet.title == "_meta":
                        continue
                    if sheet.title not in TABLES:
                        raise ExchangeError("Неизвестный лист", sheet.title)
                    if (sheet.max_column or 0) > 100 or (sheet.max_row or 0) > MAX_ROWS + 1:
                        raise ExchangeError("Превышены размеры листа", sheet.title)

                    def values(worksheet):
                        for row_number, row in enumerate(worksheet, 1):
                            for cell in row:
                                if cell.data_type in ("f", "e"):
                                    raise ExchangeError(
                                        "Формулы и ошибки Excel недопустимы",
                                        worksheet.title,
                                        row_number,
                                    )
                            yield [cell.value for cell in row]

                    result[sheet.title] = normalize_rows(sheet.title, values(sheet))
            finally:
                workbook.close()
        if not result or sum(map(len, result.values())) > MAX_TOTAL_ROWS:
            raise ExchangeError("Нет таблиц или превышен общий лимит строк")
        return result
    except ExchangeError:
        raise
    except (
        BadZipFile,
        KeyError,
        ValueError,
        OSError,
        EOFError,
        ParseError,
        DefusedXmlException,
        RuntimeError,
        NotImplementedError,
    ) as error:
        raise ExchangeError("Повреждённый файл или отсутствует manifest.json/_meta") from error


def write_csv(name: str, rows: list[dict]) -> bytes:
    columns = [c.name for c in columns_for(name)]
    if not rows:
        return (",".join(columns) + "\r\n").encode("utf-8-sig")
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(columns)
    writer.writerows([encode_cell(row.get(c)) for c in columns] for row in rows)
    return buffer.getvalue().encode("utf-8-sig")


def serialize(tables: dict[str, list[dict]], format: str) -> bytes:
    output = io.BytesIO()
    if format == "csv":
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            for name, data in [
                ("manifest.json", json.dumps({"format_version": FORMAT_VERSION}).encode())
            ] + [(name + ".csv", write_csv(name, rows)) for name, rows in tables.items()]:
                info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                archive.writestr(info, data)
    elif format == "xlsx":
        # Validate before opening write-only XML streams, so rejected exports do not
        # leave partially written worksheets or silently truncated cell contents.
        for name, rows in tables.items():
            for row_number, row in enumerate(rows, 2):
                for column in columns_for(name):
                    if len(encode_cell(row.get(column.name))) > 32767:
                        raise ExchangeError(
                            "Ячейка длиннее лимита XLSX (32767); используйте CSV",
                            name,
                            row_number,
                            column.name,
                        )
        workbook = Workbook(write_only=True)
        workbook.create_sheet("_meta").append(["format_version", FORMAT_VERSION])
        for name, rows in tables.items():
            sheet = workbook.create_sheet(name)
            columns = [c.name for c in columns_for(name)]
            sheet.append(columns)
            for row in rows:
                cells = []
                for key in columns:
                    value = encode_cell(row.get(key))
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.data_type = "s"
                    cells.append(cell)
                sheet.append(cells)
        workbook.save(output)
        workbook.close()
    else:
        raise ExchangeError("Неизвестный формат")
    return output.getvalue()
