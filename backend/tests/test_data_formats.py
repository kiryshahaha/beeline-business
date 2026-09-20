"""File contracts: no database required, every table is exercised in both formats."""

import copy
import io
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

from app.db.base import Base
from app.modules.data_exchange.formats import MAX_FILE_BYTES, ExchangeError, parse_file, serialize
from app.modules.data_exchange.registry import TABLES
from generate_synthetic import generate_dataset
from seed_synthetic import validate_database_url


class DataFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = generate_dataset(tickets=40, workers=8, days=2)

    def test_all_tables_round_trip_in_both_formats(self):
        first = parse_file(serialize(self.data, "csv"), "data.zip")
        for format, extension in (("csv", "zip"), ("xlsx", "xlsx")):
            with self.subTest(format=format):
                parsed = parse_file(serialize(self.data, format), "data." + extension)
                self.assertEqual(set(parsed), set(TABLES))
                self.assertEqual(parsed, first)
                self.assertEqual(len(parsed["tickets"]), 40)
                self.assertTrue(all(parsed[name] for name in TABLES))

    def test_synthetic_data_is_reproducible_and_varies_with_seed(self):
        self.assertEqual(self.data, generate_dataset(tickets=40, workers=8, days=2))
        self.assertNotEqual(
            self.data["locations"],
            generate_dataset(seed=43, tickets=40, workers=8, days=2)["locations"],
        )
        self.assertEqual(
            {r["transport_type"] for r in self.data["workers"]},
            {"car", "walking", "bicycle", "public_transport"},
        )
        locations = {row["id"]: row for row in self.data["locations"]}
        for ticket in self.data["tickets"]:
            if "missing_coordinates" in ticket["title"]:
                self.assertIsNone(locations[ticket["location_id"]]["latitude"])
                self.assertIsNone(locations[ticket["location_id"]]["longitude"])

    def test_all_domain_tables_are_exported_and_seeding_rejects_work_databases(self):
        self.assertEqual(
            set(Base.metadata.tables) - set(TABLES),
            {
                "refresh_tokens",
                "calendar_tokens",
                "push_subscriptions",
                "data_imports",
                "planning_plans",
                "planning_plan_routes",
            },
        )
        for url in ("postgresql://localhost/production", "sqlite:///example_test"):
            with self.subTest(url=url), self.assertRaises(ValueError):
                validate_database_url(url)
        self.assertEqual(
            validate_database_url("postgresql://localhost/example_test").database,
            "example_test",
        )

    def test_excel_long_cells_fail_explicitly_and_csv_keeps_the_entire_text(self):
        data = copy.deepcopy(self.data)
        data["tickets"][0]["description"] = "x" * 32768
        with self.assertRaises(ExchangeError) as result:
            serialize(data, "xlsx")
        self.assertEqual(result.exception.detail["table"], "tickets")
        self.assertEqual(result.exception.detail["column"], "description")
        self.assertEqual(result.exception.detail["row"], 2)
        parsed = parse_file(serialize(data, "csv"), "data.zip")
        self.assertEqual(parsed["tickets"][0]["description"], "x" * 32768)

    def test_empty_tables_and_malformed_workbook_xml(self):
        empty = {name: [] for name in TABLES}
        for format, extension in (("csv", "zip"), ("xlsx", "xlsx")):
            self.assertEqual(parse_file(serialize(empty, format), "data." + extension), empty)
        output = io.BytesIO()
        with ZipFile(io.BytesIO(serialize(self.data, "xlsx"))) as original:
            with ZipFile(output, "w", ZIP_DEFLATED) as corrupt:
                for entry in original.infolist():
                    content = original.read(entry)
                    if entry.filename == "xl/worksheets/sheet2.xml":
                        content = b"<broken"
                    corrupt.writestr(entry, content)
        with self.assertRaises(ExchangeError):
            parse_file(output.getvalue(), "data.xlsx")

    def test_formula_like_strings_and_null_are_preserved_safely(self):
        values = [
            "",
            "control\x01character",
            "\\E",
            '\\T"literal"',
            "=SUM(1,2)",
            "+cmd",
            "-7",
            "@evil",
            "'literal",
            "\\N",
            "'\\N",
            "  =1+1",
            "\t=1+1",
            "ёжик,;\nвторая строка",
        ]
        for value in values:
            for format, extension in (("csv", "zip"), ("xlsx", "xlsx")):
                with self.subTest(value=value, format=format):
                    data = {"cities": [{"id": 1, "name": value}]}
                    content = serialize(data, format)
                    self.assertEqual(parse_file(content, "data." + extension), data)
                    if format == "xlsx":
                        book = load_workbook(io.BytesIO(content))
                        self.assertEqual(book["cities"]["B2"].data_type, "s")
                        book.close()

    def test_csv_delimiters_bom_and_errors_with_coordinates(self):
        for delimiter in (",", ";", "\t"):
            self.assertEqual(
                parse_file(
                    f"\ufeffid{delimiter}name\n1{delimiter}Город\n".encode(), "cities.csv", "cities"
                )["cities"][0]["name"],
                "Город",
            )
        with self.assertRaises(ExchangeError) as result:
            parse_file(b"id,name\nnot-an-id,City\n", "data.csv", "cities")
        self.assertEqual(result.exception.detail["row"], 2)
        self.assertEqual(result.exception.detail["column"], "id")

    def test_invalid_structure_and_values_are_rejected(self):
        for content, filename, entity in (
            (b"id,name\n1,A\n1,B", "x.csv", "cities"),
            (b"id,id\n1,2", "x.csv", "cities"),
            (b"id,unknown\n1,a", "x.csv", "cities"),
            (b"id,name\n1.5,City", "x.csv", "cities"),
            (b"id,name\n2147483648,City", "x.csv", "cities"),
            (b"id,name\n1,City,Extra", "x.csv", "cities"),
            (b"id,name\n1,City", "x.csv", "users; DROP TABLE users"),
            (b"id,name\n1,City", "x.csv", None),
            (b"not a workbook", "x.xlsx", None),
            (b"not a file", "x.xls", None),
            (b"", "x.csv", "cities"),
            (b"x" * (MAX_FILE_BYTES + 1), "x.csv", "cities"),
        ):
            with self.subTest(filename=filename, entity=entity, prefix=content[:30]):
                with self.assertRaises(ExchangeError):
                    parse_file(content, filename, entity)

    def test_timezone_enum_and_coordinates_are_typed(self):
        for entity, key, invalid in (
            ("workers", "transport_type", "helicopter"),
            ("tickets", "visit_window_start", "2026-09-21T09:00:00"),
            ("locations", "latitude", "NaN"),
            ("locations", "latitude", "1e1000"),
            ("locations", "latitude", "55.1234567"),
        ):
            data = copy.deepcopy(self.data)
            data[entity][0][key] = invalid
            with self.subTest(entity=entity), self.assertRaises(ExchangeError):
                parse_file(serialize(data, "csv"), "data.zip")

    def test_excel_formulas_and_unknown_sheets_are_rejected(self):
        for mode in ("formula", "unknown", "version"):
            book = load_workbook(io.BytesIO(serialize(self.data, "xlsx")))
            if mode == "formula":
                book["cities"]["B2"] = "=1+1"
            elif mode == "unknown":
                book.create_sheet("password_hashes")
            else:
                book["_meta"]["B1"] = "99"
            output = io.BytesIO()
            book.save(output)
            book.close()
            with self.subTest(mode=mode), self.assertRaises(ExchangeError):
                parse_file(output.getvalue(), "data.xlsx")

    def test_csv_archive_unknown_paths_are_rejected_without_extraction(self):
        output = io.BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            archive.writestr("manifest.json", '{"format_version":"1"}')
            archive.writestr("../cities.csv", "id,name\n1,A")
        with self.assertRaises(ExchangeError):
            parse_file(output.getvalue(), "data.zip")
