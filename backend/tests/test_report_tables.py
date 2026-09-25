"""Spreadsheet text rule and readable report tables, without a database."""

import csv
import io
import unittest
from datetime import UTC, date, datetime
from decimal import Decimal
from enum import Enum
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from app.core.spreadsheet import starts_like_formula, xlsx_length
from app.modules.reports.tables import (
    CellTooLongError,
    Table,
    build_file,
    csv_cell,
    write_csv,
    write_csv_zip,
    write_xlsx,
)


class Color(Enum):
    RED = "red"


class SpreadsheetRuleTests(unittest.TestCase):
    def test_formula_prefixes_after_any_leading_whitespace(self):
        for value in (
            "=1+1",
            "+7 999",
            "-2+3",
            "@SUM(A1)",
            " =1",
            "  \t=1",
            "\u00a0=1",
            "\u3000@a",
            "\t",
            "\rtext",
            "\ntext",
            "=",
        ):
            with self.subTest(value=value):
                self.assertTrue(starts_like_formula(value))
        for value in ("", "text", "1-2", "a=b", "'=1", "№5", " text", "#N/A"):
            with self.subTest(value=value):
                self.assertFalse(starts_like_formula(value))

    def test_xlsx_length_counts_utf16_units(self):
        self.assertEqual(xlsx_length("abc"), 3)
        self.assertEqual(xlsx_length("ёжик"), 4)
        self.assertEqual(xlsx_length("🙂"), 2)


class ReportTableTests(unittest.TestCase):
    def test_csv_cells_escape_text_but_not_numbers(self):
        cases = {
            "=1+1": "'=1+1",
            " \t@a": "' \t@a",
            "-5": "'-5",
            "обычный текст": "обычный текст",
            "'цитата": "'цитата",
            -5: "-5",
            Decimal("-1.500000"): "-1.500000",
            2.5: "2.5",
            True: "true",
            None: "",
            date(2026, 9, 25): "2026-09-25",
            datetime(2026, 9, 25, 7, tzinfo=UTC): "2026-09-25T07:00:00+00:00",
            Color.RED: "red",
        }
        for value, expected in cases.items():
            with self.subTest(value=value):
                self.assertEqual(csv_cell(value), expected)

    def test_csv_file_has_bom_crlf_and_quotes(self):
        exported = build_file(
            lambda file: write_csv(file, Table("t", ("a", "b"), [["=x", 'в "кавычках"\n']])),
            "text/csv",
            "t.csv",
        )
        content = b"".join(exported.chunks())
        self.assertTrue(exported.file.closed)
        self.assertEqual(exported.size, len(content))
        self.assertEqual(content.decode("utf-8"), '\ufeffa,b\r\n\'=x,"в ""кавычках""\n"\r\n')

    def test_xlsx_types_are_explicit(self):
        values = [
            "=1+1",
            "#DIV/0!",
            "текст",
            "a\x00b\x1fc\ufffe",
            -3,
            Decimal("55.75"),
            False,
            None,
        ]
        columns = tuple(f"c{i}" for i in range(len(values)))
        exported = build_file(
            lambda file: write_xlsx(file, [Table("t", columns, [values])]), "x", "t.xlsx"
        )
        workbook = load_workbook(io.BytesIO(b"".join(exported.chunks())), data_only=False)
        self.addCleanup(workbook.close)
        cells = list(workbook["t"].iter_rows(min_row=2))[0]
        self.assertEqual(
            [(cell.data_type, cell.value) for cell in cells[:7]],
            [
                ("s", "=1+1"),
                ("s", "#DIV/0!"),
                ("s", "текст"),
                ("s", "a\ufffdb\ufffdc\ufffd"),
                ("n", -3),
                ("n", 55.75),
                ("b", False),
            ],
        )
        self.assertEqual([cell.quotePrefix for cell in cells[:3]], [True, True, False])
        self.assertIsNone(cells[7].value)

    def test_too_long_xlsx_text_is_refused_without_temporary_files(self):
        ok = build_file(
            lambda file: write_xlsx(file, [Table("t", ("text",), [["я" * 32_767]])]), "x", "t"
        )
        ok.file.close()
        for text in ("я" * 32_768, "🙂" * 16_384):
            with self.subTest(length=len(text)), self.assertRaises(CellTooLongError) as caught:
                build_file(
                    lambda file, text=text: write_xlsx(
                        file,
                        [
                            Table("first", ("n",), [[1]]),
                            Table("second", ("id", "text"), [[7, "ok"], [8, text]]),
                        ],
                    ),
                    "x",
                    "t",
                )
            error = caught.exception
            self.assertEqual((error.sheet, error.column, error.length), ("second", "text", 32_768))
            self.assertEqual(error.values[0], 8)
            self.assertEqual(ALL_TEMP_FILES, [])

    def test_csv_zip_holds_one_file_per_table(self):
        exported = build_file(
            lambda file: write_csv_zip(
                file, [Table("a", ("x",), [["+1"]]), Table("b", ("y",), [])]
            ),
            "application/zip",
            "t.zip",
        )
        with ZipFile(io.BytesIO(b"".join(exported.chunks()))) as archive:
            self.assertEqual(archive.namelist(), ["a.csv", "b.csv"])
            rows = list(csv.reader(io.StringIO(archive.read("a.csv").decode("utf-8-sig"))))
            self.assertEqual(rows, [["x"], ["'+1"]])
            self.assertEqual(archive.read("b.csv"), "\ufeffy\r\n".encode())


if __name__ == "__main__":
    unittest.main()
