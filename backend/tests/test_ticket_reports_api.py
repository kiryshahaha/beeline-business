"""Ticket report downloads, filtering and role checks."""

import csv
import io
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.security import create_access_token
from app.db.models import Building, City, District, Entrance, Location, Street, Ticket
from app.db.session import get_session
from app.main import app
from app.modules.data_exchange.formats import parse_file, serialize
from app.modules.data_exchange.service import import_data
from app.modules.reports import repository, service
from generate_synthetic import generate_dataset
from tests.support import DatabaseTestCase

EXPORT = "/api/v1/reports/tickets/export"
WINDOW_START = datetime(2026, 9, 25, 7, tzinfo=UTC)


def csv_rows(response) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"), newline="")))


def xlsx_rows(response) -> list[dict[str, object]]:
    """Read cells with data_only=False: a formula would keep its "f" type and text."""
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    try:
        sheet = workbook["tickets"]
        headers = [cell.value for cell in sheet[1]]
        return [
            {header: cell for header, cell in zip(headers, row, strict=True)}
            for row in sheet.iter_rows(min_row=2)
        ]
    finally:
        workbook.close()


class TicketExportApiTests(DatabaseTestCase):
    """Exercise the file contract against a real isolated PostgreSQL schema."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        dataset = generate_dataset(seed=803, tickets=32, workers=8, days=2)
        normalized = parse_file(serialize(dataset, "csv"), "dataset.zip")
        with Session(cls.engine) as session:
            result = import_data(session, normalized)
        cls.ids = result["id_map"]

    def setUp(self):
        super().setUp()

        def override_session():
            with Session(bind=self.connection, join_transaction_mode="create_savepoint") as session:
                yield session

        app.dependency_overrides[get_session] = override_session
        self.addCleanup(app.dependency_overrides.pop, get_session)
        self.client = self.enterContext(TestClient(app))
        self.observer = self.auth(1)
        self.foreman = self.auth(3)
        self.worker = self.auth(9)

    def auth(self, source_id: int) -> dict[str, str]:
        user_id = self.ids["users"][str(source_id)]
        return {"Authorization": f"Bearer {create_access_token({'sub': str(user_id)})}"}

    def test_csv_export_returns_filtered_rows_and_download_headers(self):
        response = self.client.get(
            "/api/v1/reports/tickets/export",
            params={
                "format": "csv",
                "status": "completed",
                "city_id": self.ids["cities"]["1"],
            },
            headers=self.observer,
        )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "text/csv; charset=utf-8")
        self.assertIn('attachment; filename="tickets.csv"', response.headers["content-disposition"])

        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertGreater(len(rows), 0)
        self.assertEqual(
            set(rows[0]),
            {
                "id",
                "location_id",
                "title",
                "description",
                "work_type",
                "work_type_id",
                "category",
                "priority",
                "received_at",
                "sla_deadline_at",
                "required_transport_type",
                "service_duration_source",
                "status",
                "visit_window_start",
                "visit_window_end",
                "planned_start_at",
                "planned_end_at",
                "estimated_duration_minutes",
                "actual_duration_minutes",
                "created_at",
                "assigned_worker_id",
                "is_pinned",
                "city_id",
                "city",
                "service_area_id",
                "district",
                "street_id",
                "street",
                "building_id",
                "building_number",
                "block",
                "entrance_id",
                "entrance_number",
                "floor",
                "apartment",
                "latitude",
                "longitude",
                "address",
            },
        )
        self.assertTrue(all(row["status"] == "completed" for row in rows))
        self.assertTrue(all(row["city_id"] == str(self.ids["cities"]["1"]) for row in rows))

    def test_xlsx_export_returns_workbook_and_applies_service_area_filter(self):
        response = self.client.get(
            "/api/v1/reports/tickets/export",
            params={"format": "xlsx", "service_area_id": self.ids["service_areas"]["101"]},
            headers=self.observer,
        )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            'attachment; filename="tickets.xlsx"', response.headers["content-disposition"]
        )

        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["tickets"])
        rows = list(workbook["tickets"].iter_rows(values_only=True))
        self.assertGreater(len(rows), 1)
        headers = list(rows[0])
        area_index = headers.index("service_area_id")
        self.assertTrue(
            all(row[area_index] == self.ids["service_areas"]["101"] for row in rows[1:])
        )

    def test_xlsx_is_default_format(self):
        response = self.client.get(
            "/api/v1/reports/tickets/export",
            params={"status": "planned"},
            headers=self.observer,
        )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(
            'attachment; filename="tickets.xlsx"', response.headers["content-disposition"]
        )

    def test_brigade_filter_is_applied(self):
        brigade_id = self.ids["brigades"]["1"]
        response = self.client.get(
            "/api/v1/reports/tickets/export",
            params={"format": "csv", "brigade_id": brigade_id},
            headers=self.observer,
        )

        self.assertEqual(response.status_code, 200, response.text)
        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertGreater(len(rows), 0)
        self.assertTrue(all(row["assigned_worker_id"] for row in rows))

    def test_export_requires_observer_and_valid_format(self):
        self.assertEqual(
            self.client.get("/api/v1/reports/tickets/export").status_code,
            401,
        )
        for headers in (self.foreman, self.worker):
            response = self.client.get(
                "/api/v1/reports/tickets/export",
                headers=headers,
            )
            self.assertEqual(response.status_code, 403, response.text)

        response = self.client.get(
            "/api/v1/reports/tickets/export",
            params={"format": "pdf"},
            headers=self.observer,
        )
        self.assertEqual(response.status_code, 422, response.text)


class TicketExportSafetyTests(DatabaseTestCase):
    """A30/A33: user text never runs as a formula, values and numbers are not lost."""

    def setUp(self):
        super().setUp()

        def override_session():
            with Session(bind=self.connection, join_transaction_mode="create_savepoint") as session:
                yield session

        app.dependency_overrides[get_session] = override_session
        self.addCleanup(app.dependency_overrides.pop, get_session)
        self.client = self.enterContext(TestClient(app))
        with Session(bind=self.connection, join_transaction_mode="create_savepoint") as session:
            observer_id = session.execute(
                text("""
                    INSERT INTO users (name, surname, username, password_hash, role)
                    VALUES ('Отчёт', 'Наблюдатель', :username, 'unused-hash', 'observer')
                    RETURNING id
                """),
                {"username": "report_" + uuid4().hex[:8]},
            ).scalar_one()
            session.commit()
        self.observer = {
            "Authorization": f"Bearer {create_access_token({'sub': str(observer_id)})}"
        }

    def add_place(
        self,
        *,
        city="Тестоград",
        district="Центр",
        street="ул. Тестовая",
        number="1",
        block=None,
        entrance=None,
        floor=None,
        apartment=None,
    ):
        with Session(bind=self.connection, join_transaction_mode="create_savepoint") as session:
            city_row = City(name=city)
            session.add(city_row)
            session.flush()
            district_row = District(city_id=city_row.id, name=district)
            street_row = Street(city_id=city_row.id, name=street)
            session.add_all([district_row, street_row])
            session.flush()
            building = Building(
                city_id=city_row.id,
                street_id=street_row.id,
                service_area_id=self.service_area_for_district(district_row.id),
                number=number,
                block=block,
            )
            session.add(building)
            session.flush()
            entrance_id = None
            if entrance is not None:
                entrance_row = Entrance(building_id=building.id, number=entrance)
                session.add(entrance_row)
                session.flush()
                entrance_id = entrance_row.id
            location = Location(
                building_id=building.id,
                entrance_id=entrance_id,
                floor=floor,
                apartment=apartment,
                latitude=Decimal("55.751244"),
                longitude=Decimal("-37.618423"),
            )
            session.add(location)
            session.commit()
            return city_row.id, location.id

    def add_ticket(self, location_id, **values):
        with Session(bind=self.connection, join_transaction_mode="create_savepoint") as session:
            ticket = Ticket(
                location_id=location_id,
                title=values.pop("title", "Обычная заявка"),
                description=values.pop("description", None),
                work_type=values.pop("work_type", None),
                visit_window_start=WINDOW_START,
                visit_window_end=WINDOW_START + timedelta(hours=4),
                estimated_duration_minutes=values.pop("estimated_duration_minutes", 60),
                **values,
            )
            session.add(ticket)
            session.commit()
            return ticket.id

    def export(self, **params):
        return self.client.get(EXPORT, params=params, headers=self.observer)

    def test_formula_like_text_is_literal_in_every_text_field(self):
        texts = {
            "title": "=1+1",
            "description": " \t@SUM(A1)",
            "work_type": "+7 (999) 000-00-00",
            "city": "-2+3",
            "district": "@district",
            "street": '=HYPERLINK("http://example.invalid")',
            "building_number": "+12",
            "block": "#N/A",
            "entrance_number": "-1",
            "apartment": "=A1",
        }
        city_id, location_id = self.add_place(
            city=texts["city"],
            district=texts["district"],
            street=texts["street"],
            number=texts["building_number"],
            block=texts["block"],
            entrance=texts["entrance_number"],
            floor=-2,
            apartment=texts["apartment"],
        )
        ticket_id = self.add_ticket(
            location_id,
            title=texts["title"],
            description=texts["description"],
            work_type=texts["work_type"],
            priority=2,
        )
        address = (
            '-2+3, @district, =HYPERLINK("http://example.invalid"), д. +12, #N/A, '
            "подъезд -1, этаж -2, кв./пом. =A1"
        )

        response = self.export(format="xlsx", city_id=city_id)
        self.assertEqual(response.status_code, 200, response.text)
        [row] = xlsx_rows(response)
        for column, value in {**texts, "address": address}.items():
            with self.subTest(format="xlsx", column=column):
                self.assertEqual(row[column].data_type, "s")
                self.assertEqual(row[column].value, value)
                self.assertTrue(row[column].quotePrefix)
        for column, value in (
            ("id", ticket_id),
            ("priority", 2),
            ("floor", -2),
            ("estimated_duration_minutes", 60),
        ):
            with self.subTest(format="xlsx", column=column):
                self.assertEqual(row[column].data_type, "n")
                self.assertEqual(row[column].value, value)
        self.assertEqual(row["longitude"].data_type, "n")
        self.assertAlmostEqual(row["longitude"].value, -37.618423)
        self.assertEqual(row["category"].value, "repair")
        self.assertFalse(row["category"].quotePrefix)

        response = self.export(format="csv", city_id=city_id)
        self.assertEqual(response.status_code, 200, response.text)
        [row] = csv_rows(response)
        for column, value in {**texts, "address": address}.items():
            with self.subTest(format="csv", column=column):
                expected = value if column == "block" else "'" + value
                self.assertEqual(row[column], expected)
        self.assertEqual(row["floor"], "-2")
        self.assertEqual(row["longitude"], "-37.618423")
        self.assertEqual(row["id"], str(ticket_id))
        self.assertEqual(row["category"], "repair")

    def test_line_breaks_and_control_characters_do_not_break_files(self):
        city_id, location_id = self.add_place()
        described = self.add_ticket(
            location_id,
            title="Вертикальная\x0bтабуляция\x1bи ESC",
            description='\n=1+1\r\nвторая строка; запятая, кавычка "',
        )

        response = self.export(format="xlsx", city_id=city_id)
        self.assertEqual(response.status_code, 200, response.text)
        [row] = xlsx_rows(response)
        self.assertEqual(row["id"].value, described)
        self.assertEqual(row["title"].value, "Вертикальная\ufffdтабуляция\ufffdи ESC")
        self.assertEqual(row["description"].data_type, "s")
        # XML stores a CR LF line break as LF; the text itself is unchanged.
        self.assertEqual(row["description"].value, '\n=1+1\nвторая строка; запятая, кавычка "')

        response = self.export(format="csv", city_id=city_id)
        [row] = csv_rows(response)
        self.assertEqual(row["title"], "Вертикальная\x0bтабуляция\x1bи ESC")
        self.assertEqual(row["description"], "'\n=1+1\r\nвторая строка; запятая, кавычка \"")

    def test_long_text_is_never_cut_in_xlsx(self):
        city_id, location_id = self.add_place()
        exact = self.add_ticket(location_id, description="я" * 32_767)

        response = self.export(format="xlsx", city_id=city_id)
        self.assertEqual(response.status_code, 200, response.text)
        [row] = xlsx_rows(response)
        self.assertEqual(row["description"].value, "я" * 32_767)

        for description in ("=" + "я" * 32_767, "🙂" * 16_384):
            with self.subTest(length=len(description)):
                with Session(bind=self.connection, join_transaction_mode="create_savepoint") as s:
                    s.get(Ticket, exact).description = description
                    s.commit()
                response = self.export(format="xlsx", city_id=city_id)
                self.assertEqual(response.status_code, 422, response.text)
                detail = response.json()["detail"]
                self.assertEqual(detail["code"], "xlsx_cell_too_long")
                self.assertEqual(detail["ticket_id"], exact)
                self.assertEqual(detail["column"], "description")
                self.assertEqual(detail["length"], 32_768)
                self.assertEqual(detail["max_length"], 32_767)
                self.assertIn("CSV", detail["message"])
                self.assertEqual(ALL_TEMP_FILES, [])

                response = self.export(format="csv", city_id=city_id)
                self.assertEqual(response.status_code, 200, response.text)
                expected = "'" + description if description[0] == "=" else description
                self.assertEqual(csv_rows(response)[0]["description"], expected)

    def test_empty_export_has_headers_type_length_and_filename(self):
        for format, media_type, filename in (
            ("csv", "text/csv; charset=utf-8", "tickets.csv"),
            (
                "xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "tickets.xlsx",
            ),
        ):
            with self.subTest(format=format):
                response = self.export(format=format, city_id=2_147_483_647)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.headers["content-type"], media_type)
                self.assertEqual(
                    response.headers["content-disposition"], f'attachment; filename="{filename}"'
                )
                self.assertEqual(int(response.headers["content-length"]), len(response.content))
                self.assertEqual(response.headers["cache-control"], "no-store")
                self.assertEqual(response.headers["x-content-type-options"], "nosniff")
                if format == "csv":
                    self.assertEqual(
                        response.content,
                        ("\ufeff" + ",".join(service.EXPORT_COLUMNS) + "\r\n").encode(),
                    )
                else:
                    workbook = load_workbook(BytesIO(response.content), read_only=True)
                    rows = list(workbook["tickets"].iter_rows(values_only=True))
                    workbook.close()
                    self.assertEqual(rows, [service.EXPORT_COLUMNS])

    def test_large_export_is_streamed_in_batches_and_bounded(self):
        city_id, location_id = self.add_place()
        count = repository.FETCH_BATCH_ROWS * 2 + 7
        self.connection.execute(
            text("""
                INSERT INTO tickets (
                    location_id, title, visit_window_start, visit_window_end,
                    estimated_duration_minutes
                )
                SELECT :location_id, '=' || n, :start, :end, 30
                FROM generate_series(1, :count) AS n
            """),
            {
                "location_id": location_id,
                "start": WINDOW_START,
                "end": WINDOW_START + timedelta(hours=2),
                "count": count,
            },
        )

        response = self.export(format="csv", city_id=city_id)
        self.assertEqual(response.status_code, 200, response.text)
        rows = csv_rows(response)
        self.assertEqual(len(rows), count)
        self.assertEqual([row["id"] for row in rows], sorted((row["id"] for row in rows), key=int))
        self.assertTrue(all(row["title"].startswith("'=") for row in rows))
        response = self.export(format="xlsx", city_id=city_id)
        self.assertEqual(response.status_code, 200, response.text)
        rows = xlsx_rows(response)
        self.assertEqual(len(rows), count)
        self.assertTrue(all(row["title"].data_type == "s" for row in rows))

        with patch.object(service, "MAX_EXPORT_ROWS", count - 1):
            for format in ("csv", "xlsx"):
                response = self.export(format=format, city_id=city_id)
                self.assertEqual(response.status_code, 422, response.text)
                detail = response.json()["detail"]
                self.assertEqual(detail["code"], "report_too_large")
                self.assertEqual((detail["rows"], detail["max_rows"]), (count, count - 1))
            # Rows committed after the count still cannot stretch the file past the limit.
            with patch.object(repository, "count_tickets", return_value=0):
                response = self.export(format="csv", city_id=city_id)
            self.assertEqual(response.status_code, 422, response.text)
            self.assertEqual(response.json()["detail"]["rows"], None)

    def test_rows_come_from_a_server_side_cursor(self):
        city_id, location_id = self.add_place()
        self.add_ticket(location_id)
        with Session(bind=self.connection, join_transaction_mode="create_savepoint") as session:
            rows = repository.stream_tickets(
                session,
                limit=10,
                status=None,
                city_id=city_id,
                service_area_id=None,
                brigade_id=None,
            )
            with rows:
                self.assertEqual(next(rows)["city_id"], city_id)
                # The whole result is not fetched into memory: PostgreSQL keeps an open cursor.
                open_cursors = session.connection().exec_driver_sql(
                    "SELECT count(*) FROM pg_cursors WHERE NOT is_holdable"
                )
                self.assertEqual(open_cursors.scalar_one(), 1)
