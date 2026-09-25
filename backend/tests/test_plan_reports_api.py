"""Plan report: a separate file with route order, times, workers, reasons and totals."""

import csv
import io
from datetime import datetime
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.core.security import create_access_token
from app.db.models import Ticket, User, Worker, WorkerSkillAssignment
from app.db.session import get_session
from app.main import app
from app.modules.data_exchange.formats import parse_file, serialize
from app.modules.data_exchange.service import import_data
from app.modules.planning import router as api
from planning_scenarios import NOW, generate_planning_dataset, preview_request
from tests.planning_fakes import FeasiblePlanner, provider_factory
from tests.support import CommittedDatabaseTestCase

SHEETS = ["summary", "routes", "visits", "unassigned", "excluded_workers"]


def sheet_rows(workbook, name) -> list[dict]:
    sheet = workbook[name]
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=2)]


class PlanReportApiTests(CommittedDatabaseTestCase):
    def setUp(self):
        super().setUp()
        self.data = generate_planning_dataset()
        with (
            Session(self.engine) as session,
            patch(
                "app.modules.data_exchange.service.hash_password",
                return_value="fixture-unused-hash",
            ),
        ):
            self.receipt = import_data(session, parse_file(serialize(self.data, "csv"), "data.zip"))
        self.ids = self.receipt["id_map"]
        self.payload = preview_request(self.receipt, self.data)
        self.settings = Settings(planning_enabled=True)

        def session_dependency():
            with Session(self.engine) as session:
                yield session

        overrides = {
            get_session: session_dependency,
            api.get_planning_engine: lambda: self.engine,
            api.planning_settings: lambda: self.settings,
            api.get_provider_factory: lambda: provider_factory,
            api.get_planner_client: FeasiblePlanner,
            api.get_clock: lambda: lambda: NOW,
        }
        app.dependency_overrides.update(overrides)
        self.addCleanup(lambda: [app.dependency_overrides.pop(k, None) for k in overrides])
        self.client = self.enterContext(TestClient(app))
        self.observer = self.auth(1)

    def auth(self, source_id):
        user_id = self.ids["users"][str(source_id)]
        return {"Authorization": f"Bearer {create_access_token({'sub': str(user_id)})}"}

    def preview(self):
        response = self.client.post(
            "/api/v1/planning/preview", json=self.payload, headers=self.observer
        )
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def report(self, plan_id, **params):
        return self.client.get(
            f"/api/v1/reports/plans/{plan_id}/export", params=params, headers=self.observer
        )

    def workbook(self, response):
        self.assertEqual(response.status_code, 200, response.text)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def prepare_partial_day(self):
        """Nobody has skill 1, one engineer is off the line, texts look like formulas."""
        skill_id = self.ids["worker_skills"]["1"]
        offline = self.payload["worker_ids"][-1]
        self.formula_ticket = self.payload["ticket_ids"][1]
        with self.engine.begin() as connection:
            connection.execute(
                WorkerSkillAssignment.__table__.delete().where(
                    WorkerSkillAssignment.skill_id == skill_id
                )
            )
            connection.execute(
                update(Worker).where(Worker.user_id == offline).values(is_on_line=False)
            )
            connection.execute(
                update(Ticket)
                .where(Ticket.id == self.formula_ticket)
                .values(title='=HYPERLINK("http://example.invalid","x")')
            )
            connection.execute(
                update(User)
                .where(User.id == self.payload["worker_ids"][0])
                .values(surname="@Иванов")
            )
        return offline

    def test_xlsx_report_lists_order_times_reasons_and_metrics(self):
        offline = self.prepare_partial_day()
        plan = self.preview()
        self.assertEqual(plan["outcome"], "partial")

        workbook = self.workbook(self.report(plan["plan_id"]))
        self.assertEqual(workbook.sheetnames, SHEETS)
        summary = {row["field"].value: row["value"] for row in sheet_rows(workbook, "summary")}
        self.assertEqual(summary["plan_id"].value, plan["plan_id"])
        self.assertEqual(summary["state"].value, "ready")
        self.assertEqual(summary["outcome"].value, "partial")
        self.assertEqual(summary["route_date"].value, plan["route_date"])
        self.assertEqual(summary["objective_order"].value, "unassigned_total > travel_minutes")
        for name, value in plan["metrics"].items():
            if name == "unassigned_by_category":
                continue
            with self.subTest(metric=name):
                self.assertEqual(summary[f"metrics.{name}"].data_type, "n")
                self.assertEqual(summary[f"metrics.{name}"].value, value)
        for category, count in plan["metrics"]["unassigned_by_category"].items():
            self.assertEqual(summary[f"unassigned_by_category.{category}"].value, count)
        self.assertEqual(summary["unassigned_by_category.skill"].value, 8)
        self.assertIsNone(summary["applied_at"].value)

        visits = sheet_rows(workbook, "visits")
        expected = [
            (route["worker_id"], stop["sequence"], stop["ticket_id"])
            for route in plan["routes"]
            for stop in sorted(route["stops"], key=lambda item: item["sequence"])
        ]
        self.assertEqual(
            [(r["worker_id"].value, r["sequence"].value, r["ticket_id"].value) for r in visits],
            expected,
        )
        stops = {stop["ticket_id"]: stop for route in plan["routes"] for stop in route["stops"]}
        titles = {t["id"]: t["title"] for t in self.data["tickets"]}
        source_ids = {v: int(k) for k, v in self.ids["tickets"].items()}
        for row in visits:
            ticket_id = row["ticket_id"].value
            stop = stops[ticket_id]
            with self.subTest(ticket=ticket_id):
                for column in ("arrival_at", "service_start_at", "service_end_at"):
                    self.assertEqual(
                        datetime.fromisoformat(row[column].value),
                        datetime.fromisoformat(stop[column]),
                    )
                    self.assertTrue(row[column].value.endswith("+03:00"))
                self.assertEqual(row["service_minutes"].value, stop["effective_service_minutes"])
                self.assertIn(
                    row["category"].value, {"emergency", "connection", "repair", "additional"}
                )
                self.assertEqual(row["priority"].data_type, "n")
                self.assertTrue(row["factors"].value)
                if ticket_id != self.formula_ticket:
                    self.assertEqual(row["title"].value, titles[source_ids[ticket_id]])
        [formula] = [row for row in visits if row["ticket_id"].value == self.formula_ticket]
        self.assertEqual(formula["title"].data_type, "s")
        self.assertEqual(formula["title"].value, '=HYPERLINK("http://example.invalid","x")')
        self.assertTrue(formula["title"].quotePrefix)
        first = self.payload["worker_ids"][0]
        named = [row for row in visits if row["worker_id"].value == first]
        self.assertTrue(named)
        self.assertTrue(all(row["worker"].value.startswith("@Иванов") for row in named))
        self.assertTrue(all(row["worker"].data_type == "s" for row in named))

        routes = sheet_rows(workbook, "routes")
        self.assertEqual(
            [(r["worker_id"].value, r["visits"].value) for r in routes],
            [(route["worker_id"], len(route["stops"])) for route in plan["routes"]],
        )
        self.assertTrue(all(row["route_number"].value is None for row in routes))

        unassigned = sheet_rows(workbook, "unassigned")
        self.assertEqual(
            [row["ticket_id"].value for row in unassigned],
            [item["ticket_id"] for item in plan["unassigned"]],
        )
        self.assertIn("missing_skill", {row["reason_code"].value for row in unassigned})
        for row, item in zip(unassigned, plan["unassigned"], strict=True):
            self.assertEqual(row["reason_code"].value, item["reason"]["code"])
            self.assertEqual(row["reason_category"].value, item["reason"]["category"])
            self.assertEqual(row["reason"].value, item["reason"]["message"])
            self.assertEqual(row["candidates"].value.count(";") + 1, len(item["candidates"]))
        [excluded] = sheet_rows(workbook, "excluded_workers")
        self.assertEqual(excluded["worker_id"].value, offline)
        self.assertEqual(excluded["reason_code"].value, "worker_offline")
        self.assertEqual(excluded["reason"].value, "Инженер снят с линии")

    def test_applied_plan_report_shows_route_numbers_and_current_state(self):
        plan = self.preview()
        applied = self.client.post(
            f"/api/v1/planning/plans/{plan['plan_id']}/apply", headers=self.observer
        )
        self.assertEqual(applied.status_code, 200, applied.text)
        numbers = {r["worker_id"]: r["route_number"] for r in applied.json()["routes"]}

        workbook = self.workbook(self.report(plan["plan_id"], format="xlsx"))
        summary = {row["field"].value: row["value"] for row in sheet_rows(workbook, "summary")}
        self.assertEqual(summary["state"].value, "applied")
        self.assertIs(summary["is_current"].value, True)
        self.assertTrue(summary["applied_at"].value.endswith("+03:00"))
        routes = sheet_rows(workbook, "routes")
        self.assertEqual({r["worker_id"].value: r["route_number"].value for r in routes}, numbers)
        self.assertEqual(sheet_rows(workbook, "unassigned"), [])

        # A later change of an assigned ticket makes the applied plan historical.
        with self.engine.begin() as connection:
            connection.execute(
                update(Ticket)
                .where(Ticket.id == self.payload["ticket_ids"][0])
                .values(description="изменено после применения")
            )
        workbook = self.workbook(self.report(plan["plan_id"]))
        summary = {row["field"].value: row["value"] for row in sheet_rows(workbook, "summary")}
        self.assertIs(summary["is_current"].value, False)

    def test_csv_report_is_a_zip_of_safe_tables(self):
        self.prepare_partial_day()
        plan = self.preview()

        response = self.report(plan["plan_id"], format="csv")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "application/zip")
        self.assertEqual(
            response.headers["content-disposition"],
            f'attachment; filename="plan-{plan["plan_id"]}.zip"',
        )
        self.assertEqual(int(response.headers["content-length"]), len(response.content))
        with ZipFile(BytesIO(response.content)) as archive:
            self.assertEqual(archive.namelist(), [name + ".csv" for name in SHEETS])
            tables = {
                name: list(
                    csv.DictReader(
                        io.StringIO(archive.read(name + ".csv").decode("utf-8-sig"), newline="")
                    )
                )
                for name in SHEETS
            }
        [formula] = [r for r in tables["visits"] if r["ticket_id"] == str(self.formula_ticket)]
        self.assertEqual(formula["title"], '\'=HYPERLINK("http://example.invalid","x")')
        self.assertEqual(len(tables["visits"]), plan["metrics"]["assigned_tickets"])
        first = str(self.payload["worker_ids"][0])
        named = [row["worker"] for row in tables["visits"] if row["worker_id"] == first]
        self.assertTrue(named)
        self.assertTrue(all(name.startswith("'@Иванов") for name in named))
        self.assertEqual(
            [row["reason_code"] for row in tables["unassigned"]],
            [item["reason"]["code"] for item in plan["unassigned"]],
        )
        summary = {row["field"]: row["value"] for row in tables["summary"]}
        self.assertEqual(
            summary["metrics.assigned_tickets"], str(plan["metrics"]["assigned_tickets"])
        )
        self.assertEqual(summary["is_current"], "")

    def test_unknown_plan_roles_and_format_are_checked(self):
        plan = self.preview()
        missing = self.report(uuid4())
        self.assertEqual(missing.status_code, 404, missing.text)
        self.assertEqual(missing.json()["detail"]["code"], "plan_not_found")
        self.assertEqual(self.report(plan["plan_id"], format="pdf").status_code, 422)
        path = f"/api/v1/reports/plans/{plan['plan_id']}/export"
        self.assertEqual(self.client.get(path).status_code, 401)
        for source_id in (3, 9):  # foreman, worker
            with self.subTest(source_id=source_id):
                response = self.client.get(path, headers=self.auth(source_id))
                self.assertEqual(response.status_code, 403, response.text)

    def test_plan_report_does_not_change_the_plan_or_ticket_export(self):
        plan = self.preview()
        before = self.client.get(f"/api/v1/planning/plans/{plan['plan_id']}", headers=self.observer)
        self.assertEqual(self.report(plan["plan_id"]).status_code, 200)
        after = self.client.get(f"/api/v1/planning/plans/{plan['plan_id']}", headers=self.observer)
        self.assertEqual(before.json(), after.json())
        tickets = self.client.get(
            "/api/v1/reports/tickets/export", params={"format": "xlsx"}, headers=self.observer
        )
        workbook = self.workbook(tickets)
        self.assertEqual(workbook.sheetnames, ["tickets"])
