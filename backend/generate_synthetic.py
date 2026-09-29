"""Generate deterministic synthetic exchange packages on real Moscow addresses, without a DB.

The package is the state of the case on the morning of its last day: the three areas of
the organizer with their offices and brigades, engineers with different skills, shifts
and transport, the history of the previous days (completed and cancelled visits) and the
open demand of the last day. Routes, day plans and execution journals are not included:
the planner and the engineers' app create them.
"""

import argparse
import csv
import hashlib
import io
import json
import random
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from app.modules.data_exchange.formats import encode_cell, json_default, parse_file, serialize
from app.modules.data_exchange.registry import TABLES
from synthetic_moscow.catalog import (
    APPLIANCES,
    CANCEL_REASONS,
    DISPATCHER_NOTES,
    FOREMAN_NOTES,
    INCIDENTS,
    KIT_RESERVES,
    ORDERED_DEVICE,
    REQUEST_KINDS,
    REQUIRED_APPLIANCES,
    SKILLS,
    SPEEDS,
    TKD_PLACES,
    WORK_TYPES,
    WORKER_REPORTS,
    WorkTypeSpec,
)
from synthetic_moscow.geography import AREAS, MOSCOW, Building, district_buildings, split_zones
from synthetic_moscow.people import PeopleFactory

TZ = timezone(timedelta(hours=3))
SCENARIOS = (
    "regular_demand",
    "overlapping_windows",
    "whole_day_emergency",
    "night_emergency",
    "missing_coordinates",
    "equipment_shortage",
    "closed_history",
)
# Visit windows of the organizer's day files and how often each one occurs.
WINDOWS = ((10, 16), (12, 17), (14, 12), (16, 10), (18, 11), (20, 7))
OBSERVERS = 2
# A week of history before the planning day Monday 2026-09-21 (the last day of the period).
START_DATE = date(2026, 9, 15)
FOREMEN = 2 * len(AREAS)
FIRST_WORKER_ID = OBSERVERS + FOREMEN + 1


@dataclass(frozen=True)
class Shift:
    start: time
    end: time
    schedule: str
    night: bool = False


DAY_12 = Shift(time(10), time(22), "2/2")
EARLY_12 = Shift(time(9), time(21), "2/2")
OFFICE_8 = Shift(time(9), time(18), "5/2")
EVENING_8 = Shift(time(14), time(22), "5/2")
NIGHT_12 = Shift(time(22), time(10), "2/2", night=True)


@dataclass(frozen=True, kw_only=True)
class Layout:
    """How a compact fixture arranges the same realistic content."""

    # Visit i goes to the brigade of engineer i % engineers.
    by_worker: bool = False
    # Visit types in turn instead of the 40/40/10/10 mix.
    categories: tuple[str, ...] | None = None
    # One window (local times, may cross midnight) for every visit of the day.
    window: tuple[time, time] | None = None
    # Two-hour windows 10:00-22:00 in turn for the visits of one engineer.
    slots: bool = False
    shift: Shift | None = None
    # Every visit in one entrance of one building: identical coordinates.
    one_entrance: bool = False
    # Data-quality cases of the day: no coordinates, no free router, shared slots.
    scenarios: bool = True


DEFAULT_LAYOUT = Layout()

# Engineer profiles: case skills, possible extra skills, transport shares.
ROLES = {
    "universal": (
        ("Аварийные работы", "Локальные работы", "Работы на подключение и дозаказы"),
        {"Монтаж и сварка ВОЛС": 0.5, "Допуск по электробезопасности (III группа)": 0.7},
        (("car", 0.85), ("public_transport", 0.15)),
    ),
    "installer": (
        ("Работы на подключение и дозаказы",),
        {"Локальные работы": 0.4, "Настройка IPTV и видеонаблюдения": 0.4},
        (("public_transport", 0.45), ("car", 0.4), ("walking", 0.1), ("bicycle", 0.05)),
    ),
    "technician": (
        ("Локальные работы",),
        {"Работы на подключение и дозаказы": 0.35, "Настройка IPTV и видеонаблюдения": 0.5},
        (("public_transport", 0.5), ("car", 0.3), ("walking", 0.1), ("bicycle", 0.1)),
    ),
    "line": (
        ("Аварийные работы", "Локальные работы"),
        {"Монтаж и сварка ВОЛС": 0.8, "Допуск по электробезопасности (III группа)": 0.9},
        (("car", 0.9), ("public_transport", 0.1)),
    ),
}
# The first member of a brigade is always a senior engineer, so every brigade has all
# case skills; the order then mixes profiles like a real brigade roster.
ROLE_ORDER = (
    "universal",
    "installer",
    "technician",
    "line",
    "installer",
    "technician",
    "installer",
    "technician",
    "universal",
    "installer",
    "technician",
    "line",
)


@dataclass
class Package:
    planning_day: date
    tables: dict = field(default_factory=lambda: {name: [] for name in TABLES})
    scenarios: dict = field(default_factory=lambda: defaultdict(list))


def _weighted(rng: random.Random, pairs):
    total = sum(weight for _, weight in pairs)
    point = rng.random() * total
    for value, weight in pairs:
        point -= weight
        if point < 0:
            return value
    return pairs[-1][0]


def _at(day: date, hour: int, minute: int = 0) -> datetime:
    return datetime.combine(day, time(hour, minute), TZ)


class Generator:
    def __init__(
        self, *, seed, start_date, tickets, workers, days, area_codes=None, layout=DEFAULT_LAYOUT
    ):
        self.rng = random.Random(seed)
        self.layout = layout
        self.all_on_duty = layout.shift is not None
        self.area_specs = [a for a in AREAS if area_codes is None or a.code in area_codes]
        if not self.area_specs or FOREMEN % len(self.area_specs):
            raise ValueError("Нужен один или три участка кейса")
        self.brigades_per_area = FOREMEN // len(self.area_specs)
        self.seed = seed
        self.start_date = start_date
        self.days = days
        self.ticket_count = tickets
        self.worker_count = workers
        self.today = start_date + timedelta(days=days - 1)
        # Reference records exist long before the period; the snapshot is 08:00 today.
        self.stamp = _at(start_date - timedelta(days=90), 9)
        self.snapshot = _at(self.today, 8)
        self.package = Package(self.today)
        self.tables = self.package.tables
        self.people = PeopleFactory(self.rng)
        self.streets: dict[tuple[int, str], int] = {}
        self.buildings: dict[tuple[int, str, str, str | None], int] = {}
        self.entrances: dict[tuple[int, str], int] = {}
        self.locations: dict[tuple[int, int | None, str | None], int] = {}
        self.neighbours: dict[int, tuple] = {}

    # -- rows -------------------------------------------------------------------------------

    def add(self, entity, **values):
        for column in TABLES[entity].columns:
            key = column.name
            if key in values or key == "password_hash":
                continue
            if key == "id":
                values[key] = len(self.tables[entity]) + 1
            elif key in ("created_at", "updated_at", "assigned_at", "next_attempt_at"):
                values[key] = self.stamp
            elif column.default is not None and getattr(column.default, "arg", None) is not None:
                arg = column.default.arg
                values[key] = arg() if callable(arg) else arg
            elif column.nullable:
                values[key] = None
        self.tables[entity].append(values)
        return values

    def street_id(self, city_id: int, name: str) -> int:
        key = (city_id, name)
        if key not in self.streets:
            self.streets[key] = self.add("streets", city_id=city_id, name=name)["id"]
        return self.streets[key]

    def building_id(self, *, city_id, area_id, street, number, block) -> int:
        key = (city_id, street, number, block)
        if key not in self.buildings:
            self.buildings[key] = self.add(
                "buildings",
                city_id=city_id,
                service_area_id=area_id,
                street_id=self.street_id(city_id, street),
                number=number,
                block=block,
            )["id"]
        return self.buildings[key]

    def entrance_id(self, building_id: int, number: str) -> int:
        key = (building_id, number)
        if key not in self.entrances:
            self.entrances[key] = self.add("entrances", building_id=building_id, number=number)[
                "id"
            ]
        return self.entrances[key]

    def location_id(self, *, building_id, entrance_id, apartment, floor, point) -> int:
        key = (building_id, entrance_id, apartment)
        if key not in self.locations:
            latitude, longitude = point if point is not None else (None, None)
            self.locations[key] = self.add(
                "locations",
                building_id=building_id,
                entrance_id=entrance_id,
                floor=floor,
                apartment=apartment,
                latitude=latitude,
                longitude=longitude,
            )["id"]
        return self.locations[key]

    # -- reference ---------------------------------------------------------------------------

    def reference(self):
        self.add("cities", id=1, name=MOSCOW)
        self.areas = []
        for index, area in enumerate(self.area_specs, 1):
            area_id = 100 + index
            self.add("districts", id=index, city_id=1, name=area.name)
            self.add(
                "service_areas",
                id=area_id,
                code=f"district_{index}",
                name=area.name,
                description="Районы Москвы: " + ", ".join(d.name for d in area.districts),
            )
            self.add("divisions", id=index, service_area_id=area_id)
            # The office row follows the requests, so location 1 is the first request.
            office_id = index
            self.areas.append({"area": area, "id": area_id, "office_id": office_id})
        for index, skill in enumerate(SKILLS, 1):
            self.add("worker_skills", id=index, skill=skill)
        self.skill_ids = {skill: index for index, skill in enumerate(SKILLS, 1)}
        for index, spec in enumerate(APPLIANCES, 1):
            self.add(
                "appliances",
                id=index,
                name=spec.name,
                type=spec.type,
                description=spec.description,
                unit=spec.unit,
                is_active=spec.is_active,
            )
        self.appliance_ids = {spec.key: index for index, spec in enumerate(APPLIANCES, 1)}
        self.work_types = {}
        for index, spec in enumerate(WORK_TYPES, 1):
            self.add(
                "work_types",
                id=index,
                name=spec.name,
                code=spec.code,
                category=spec.category,
                default_priority=spec.priority,
                travel_minutes=spec.travel_minutes,
                work_minutes=spec.work_minutes,
                documents_minutes=spec.documents_minutes,
                norm_minutes=spec.travel_minutes + spec.work_minutes + spec.documents_minutes,
            )
            self.add(
                "work_type_planning_rules",
                work_type_id=index,
                service_duration_source="work_norm",
                configured_by=1,
            )
            self.add(
                "work_type_required_skills",
                work_type_id=index,
                skill_id=self.skill_ids[spec.skill],
            )
            for key, quantity in REQUIRED_APPLIANCES.get(spec.code, ()):
                self.add(
                    "work_type_required_appliances",
                    work_type_id=index,
                    appliance_id=self.appliance_ids[key],
                    quantity=quantity,
                )
            self.work_types[spec.code] = (index, spec)

    # -- people ------------------------------------------------------------------------------

    def user(self, role: str, person) -> int:
        return self.add(
            "users",
            name=person.name,
            surname=person.surname,
            lastname=person.lastname,
            username=person.username,
            role=role,
        )["id"]

    def staff(self):
        for index in range(OBSERVERS):
            self.user("observer", self.people.person(female=index == 0))
        self.brigades = []
        foremen = [self.people.person(distinct_surname=True) for _ in range(FOREMEN)]
        foreman_ids = [self.user("foreman", person) for person in foremen]
        per_area = self.brigades_per_area
        for index, (person, foreman_id) in enumerate(zip(foremen, foreman_ids, strict=True)):
            area = self.areas[index // per_area]
            zone = split_zones(area["area"], per_area)[index % per_area]
            brigade = self.add(
                "brigades",
                id=index + 1,
                name=f"Бригада {person.short_name}",
                foreman_id=foreman_id,
                office_id=area["office_id"],
                division_id=index // per_area + 1,
            )
            self.brigades.append(
                {
                    "id": brigade["id"],
                    "area": area,
                    "zone": zone,
                    "members": [],
                    "buildings": [b for d in zone for b in district_buildings(d.name)],
                }
            )
        self.workers = {}
        for number in range(self.worker_count):
            # Alternate areas first, then brigades inside an area.
            count = len(self.areas)
            brigade = self.brigades[(number % count) * per_area + (number // count) % per_area]
            position = len(brigade["members"])
            role = ROLE_ORDER[position % len(ROLE_ORDER)]
            self.add_worker(
                brigade,
                role,
                self.layout.shift or self.shift(role, position),
                _weighted(self.rng, ROLES[role][2]),
            )
        self.cover_transports()

    def add_worker(self, brigade, role, shift: Shift, transport, *, start_location_id=None):
        position = len(brigade["members"])
        person = self.people.person()
        user_id = self.user("worker", person)
        case_skills, extras, _ = ROLES[role]
        skills = set(case_skills)
        skills.update(skill for skill, share in extras.items() if self.rng.random() < share)
        workdays = None
        cycle_start = None
        if shift.schedule == "2/2":
            # Two days on, two off, staggered: about half of the 2/2 staff is off on any day.
            # A fixture leaves the cycle unset, so its crew works on every date it is used.
            if not self.all_on_duty:
                cycle_start = self.today - timedelta(days=(user_id * 3) % 4)
        else:
            workdays = [0, 1, 2, 3, 4] if position % 3 or self.all_on_duty else [1, 2, 3, 4, 5]
        self.add(
            "workers",
            user_id=user_id,
            workshift_start=shift.start,
            workshift_end=shift.end,
            transport_type=transport,
            is_on_line=True,
            service_area_id=brigade["area"]["id"],
            start_location_id=start_location_id,
            stock_office_id=brigade["area"]["office_id"],
            schedule_type=shift.schedule,
            cycle_start_date=cycle_start,
            workdays_mask=workdays,
        )
        for skill in SKILLS:
            if skill in skills:
                self.add(
                    "worker_skill_assignments", worker_id=user_id, skill_id=self.skill_ids[skill]
                )
        self.add("brigade_members", brigade_id=brigade["id"], worker_id=user_id)
        brigade["members"].append(user_id)
        self.workers[user_id] = {
            "person": person,
            "skills": skills,
            "shift": shift,
            "role": role,
            "brigade": brigade,
            "cycle_start": cycle_start,
            "workdays": workdays,
        }
        return user_id

    def shift(self, role: str, position: int) -> Shift:
        if role in ("line", "universal") and position % 12 == 3:
            return NIGHT_12
        return (DAY_12, DAY_12, EARLY_12, DAY_12, OFFICE_8, DAY_12, EVENING_8, EARLY_12)[
            position % 8
        ]

    def cover_transports(self):
        """Keep all four transport profiles present, as the exchange contract requires."""
        rows = [row for row in self.tables["workers"]]
        if len(rows) < 4:
            return
        juniors = [
            r for r in rows if self.workers[r["user_id"]]["role"] in ("installer", "technician")
        ]
        for transport in ("car", "public_transport", "walking", "bicycle"):
            counts = defaultdict(int)
            for row in rows:
                counts[row["transport_type"]] += 1
            if counts[transport]:
                continue
            spare = [r for r in (juniors or rows) if counts[r["transport_type"]] > 1]
            if spare:
                spare[-1]["transport_type"] = transport

    # -- demand ------------------------------------------------------------------------------

    def category_deck(self):
        deck = []
        for spec in WORK_TYPES:
            deck.extend([spec.code] * round(spec.share * 10))
        return deck

    def demand(self):
        per_day = [self.ticket_count // self.days] * self.days
        for index in range(self.ticket_count % self.days):
            per_day[-1 - index] += 1
        deck = self.category_deck()
        number = 0
        for day_index, count in enumerate(per_day):
            day = self.start_date + timedelta(days=day_index)
            today = day == self.today
            self.rng.shuffle(deck)
            for position in range(count):
                code = deck[position % len(deck)]
                if position and position % len(deck) == 0:
                    self.rng.shuffle(deck)
                if self.layout.categories:
                    code = self.layout.categories[number % len(self.layout.categories)]
                self.ticket(number, day, today, code)
                number += 1
        if self.layout.scenarios:
            self.shortage()
            self.missing_coordinates()
        else:
            self.shortage_office = None

    def pick_brigade(self, number: int, today: bool):
        if self.layout.by_worker:
            user_id = list(self.workers)[number % len(self.workers)]
            return self.workers[user_id]["brigade"]
        area = self.areas[number % len(self.areas)]
        brigades = [b for b in self.brigades if b["area"] is area and b["buildings"]]
        staffed = [b for b in brigades if b["members"]]
        return self.rng.choice(staffed or brigades)

    def window(self, number: int, day: date, spec: WorkTypeSpec, night: bool):
        if self.layout.window:
            start = datetime.combine(day, self.layout.window[0], TZ)
            end = datetime.combine(day, self.layout.window[1], TZ)
            return start, end if end > start else end + timedelta(days=1)
        if self.layout.slots:
            # Every visit of an engineer in its own two-hour slot, outages included, as
            # «Авария 20:00-22:00» in the organizer's files.
            hour = 10 + 2 * ((number // len(self.workers)) % 6)
            return _at(day, hour), _at(day, hour + 2)
        if spec.category == "emergency" and (night or self.rng.random() < 0.8):
            return _at(day, 0, 1), _at(day, 23, 59)
        hour = _weighted(self.rng, WINDOWS)
        return _at(day, hour), _at(day, hour + 2)

    def focus(self, brigade):
        """The entrance with the most apartments of the brigade's districts."""
        if not hasattr(self, "focused"):
            self.focused = max(
                (
                    (entrance.last_apartment - entrance.first_apartment, building, entrance)
                    for building in brigade["buildings"]
                    for entrance in building.entrances
                ),
                key=lambda item: (item[0], item[1].osm, item[2].number),
            )[1:]
        return self.focused

    def ticket(self, number: int, day: date, today: bool, code: str):
        rng = self.rng
        spec = self.work_types[code][1]
        brigade = self.pick_brigade(number, today)
        building: Building = rng.choice(brigade["buildings"])
        emergency = spec.category == "emergency"
        night = emergency and not today and rng.random() < 0.2
        start, end = self.window(number, day, spec, night)
        neighbour = self.neighbours.get(brigade["id"]) if self.layout.scenarios else None
        if today and not emergency and neighbour and rng.random() < 0.05:
            # Neighbours in one house asked for the same slot, e.g. after an outage.
            building, start, end = neighbour
        if emergency:
            if night:
                received = _at(day, rng.choice((0, 1, 2, 3, 4, 22, 23)), rng.randrange(0, 60))
            elif today:
                received = _at(day, 6) + timedelta(minutes=rng.randrange(0, 115))
            elif (end - start) < timedelta(hours=3):
                received = start - timedelta(minutes=rng.randrange(20, 150))
            else:
                received = _at(day, rng.randrange(7, 20), rng.randrange(0, 60))
            if end - start >= timedelta(hours=3):
                # A whole-day window opens with the outage itself.
                received = max(received, start)
        else:
            previous = day - timedelta(days=rng.choice((1, 1, 1, 2, 3)))
            received = _at(previous, rng.randrange(9, 21), rng.randrange(0, 60))
            if today:
                self.neighbours[brigade["id"]] = (building, start, end)
        entrance, apartment, floor = building.apartment(rng.randrange(10_000))
        if self.layout.one_entrance:
            building, entrance = self.focus(brigade)
            size = entrance.last_apartment - entrance.first_apartment + 1
            apartment = entrance.first_apartment + number % size
            floor = min(
                building.levels,
                (apartment - entrance.first_apartment) // max(1, -(-size // building.levels)) + 1,
            )
        if emergency and not self.layout.one_entrance:
            entrance = rng.choice(building.entrances)
        ticket = self.create_ticket(
            spec,
            brigade,
            building,
            entrance,
            apartment,
            floor,
            start,
            end,
            received,
        )
        scenarios = self.package.scenarios
        if emergency:
            scenarios["night_emergency" if night else "whole_day_emergency"].append(ticket["id"])
        if today:
            scenarios["regular_demand"].append(ticket["id"])
        else:
            self.close(ticket, spec, brigade)
            scenarios["closed_history"].append(ticket["id"])
        self.allocate(ticket, spec, brigade, ticket.pop("_appliances"))
        self.comments(ticket, spec)

    def place(self, building: Building, area_id: int, entrance, apartment, floor, city_id=1):
        building_id = self.building_id(
            city_id=city_id,
            area_id=area_id,
            street=building.street,
            number=building.number,
            block=building.block,
        )
        return self.location_id(
            building_id=building_id,
            entrance_id=self.entrance_id(building_id, entrance.number),
            apartment=None if apartment is None else str(apartment),
            floor=floor,
            point=(entrance.latitude, entrance.longitude),
        )

    def create_ticket(
        self,
        spec: WorkTypeSpec,
        brigade,
        building: Building,
        entrance,
        apartment,
        floor,
        start: datetime,
        end: datetime,
        received: datetime,
        *,
        city_id=1,
        kind=None,
    ) -> dict:
        """An open request of the day at a real apartment (or the house node for outages)."""
        emergency = spec.category == "emergency"
        if emergency:
            apartment = floor = None
        area_id = brigade["area"]["id"]
        location_id = self.place(building, area_id, entrance, apartment, floor, city_id)
        title, description, appliances, hd_type, subscribers = self.texts(
            spec, building, entrance.number, kind=kind
        )
        ticket = self.add(
            "tickets",
            location_id=location_id,
            service_area_id=area_id,
            brigade_id=brigade["id"],
            title=title,
            description=description,
            work_type=spec.name,
            work_type_id=self.work_types[spec.code][0],
            category=spec.category,
            priority=spec.priority,
            received_at=received,
            sla_deadline_at=received + timedelta(hours=24) if emergency else None,
            response_deadline_at=(
                received + timedelta(minutes=60 if subscribers >= 100 else 120)
                if emergency
                else None
            ),
            intake_source=None,
            request_type_hd=hd_type,
            required_transport_type=None,
            service_duration_source="work_norm",
            status="planned",
            lifecycle_state="waiting_assignment",
            assigned_worker_id=None,
            is_pinned=False,
            visit_window_start=start,
            visit_window_end=end,
            planned_start_at=None,
            planned_end_at=None,
            estimated_duration_minutes=spec.service_minutes,
            actual_duration_minutes=None,
            revision=1,
            execution_cycle=1,
            actual_started_at=None,
            actual_completed_at=None,
            cancel_reason=None,
            last_event_id=None,
            created_at=received,
            updated_at=received,
        )
        ticket["_appliances"] = appliances
        return ticket

    def texts(self, spec: WorkTypeSpec, building: Building, entrance: str, *, kind=None):
        rng = self.rng
        values = {
            "speed": rng.choice(SPEEDS),
            "meters": rng.randrange(8, 31),
            "entrance": entrance,
            "floor": building.levels,
            "subscribers": 0,
        }
        if spec.category == "emergency":
            incident = rng.choice(INCIDENTS)
            subscribers = min(building.flats, rng.randrange(12, 160))
            values["subscribers"] = subscribers
            values["place"] = rng.choice(TKD_PLACES).format(**values)
            values["time"] = f"{rng.randrange(0, 8):02d}:{rng.randrange(0, 60):02d}"
            return (
                incident.title,
                incident.description.format(**values),
                incident.appliances,
                "Авария",
                subscribers,
            )
        kinds = REQUEST_KINDS[spec.code]
        kind = kind or _weighted(rng, [(k, k.weight) for k in kinds])
        title = rng.choice(kind.titles)
        description = rng.choice(kind.descriptions).format(**values)
        appliances = kind.appliances
        if title in ORDERED_DEVICE:
            appliances = ((ORDERED_DEVICE[title], 1, 1),) + appliances
        return title, description, appliances, kind.hd_type, 0

    @staticmethod
    def works_on(worker: dict, day: date) -> bool:
        if worker["cycle_start"] is not None:
            return (day - worker["cycle_start"]).days % 4 in (0, 1)
        return day.weekday() in worker["workdays"]

    def on_duty(self, worker: dict, moment: datetime) -> bool:
        """The engineer's shift, on a working day of his schedule, covers the moment."""
        shift = worker["shift"]
        clock = moment.time()
        if shift.night:
            if clock >= shift.start:
                return self.works_on(worker, moment.date())
            # The tail of a night duty that started the evening before.
            return clock < shift.end and self.works_on(worker, moment.date() - timedelta(days=1))
        return shift.start <= clock < shift.end and self.works_on(worker, moment.date())

    def eligible(self, brigade, spec: WorkTypeSpec, start: datetime, end: datetime):
        """Qualified members of the brigade on duty at some point of the interval."""
        points = [start + (end - start) * step / 8 for step in range(9)]
        return [
            user_id
            for user_id in brigade["members"]
            if spec.skill in self.workers[user_id]["skills"]
            and any(self.on_duty(self.workers[user_id], point) for point in points)
        ]

    def close(self, ticket: dict, spec: WorkTypeSpec, brigade):
        rng = self.rng
        start, end = ticket["visit_window_start"], ticket["visit_window_end"]
        emergency = spec.category == "emergency"
        if emergency:
            # The crew on duty when the outage was reported takes it.
            start = ticket["received_at"]
            end = start + timedelta(hours=2)
        candidates = self.eligible(brigade, spec, start, end)
        cancel = rng.random() < 0.08 or not candidates
        reason = rng.choice(CANCEL_REASONS) if cancel else None
        if reason is None and not candidates:
            reason = "Нет свободного исполнителя с нужным допуском"
        worker_id = rng.choice(candidates) if candidates else None
        if cancel:
            on_site = reason == "Клиент не открыл дверь, телефон недоступен" and worker_id
            closed_at = (
                (start + timedelta(minutes=rng.randrange(10, 90)))
                if on_site
                else (ticket["received_at"] + timedelta(hours=rng.randrange(1, 12)))
            )
            ticket.update(
                status="wont_fix",
                lifecycle_state="cancelled",
                assigned_worker_id=worker_id if on_site else None,
                cancel_reason=reason,
                revision=3 if on_site else 2,
                updated_at=max(closed_at, ticket["received_at"] + timedelta(minutes=5)),
            )
            if on_site:
                self.notify(ticket, worker_id)
            return
        duration = max(10, round(spec.service_minutes * rng.uniform(0.75, 1.3)))
        worker = self.workers[worker_id]
        if emergency:
            began = ticket["received_at"] + timedelta(minutes=rng.randrange(35, 115))
            began = max(began, ticket["visit_window_start"])
        else:
            moments = [
                start + timedelta(minutes=minute)
                for minute in range(0, int((end - start).total_seconds() // 60) - 14, 5)
            ]
            fitting = [
                m
                for m in moments
                if self.on_duty(worker, m) and self.on_duty(worker, m + timedelta(minutes=duration))
            ]
            began = rng.choice(fitting or moments or [start])
            if rng.random() < 0.05:
                # A late arrival after the promised window, as in «Просрочена».
                began = end + timedelta(minutes=rng.randrange(10, 50))
        ticket.update(
            status="completed",
            lifecycle_state="completed",
            assigned_worker_id=worker_id,
            actual_started_at=began,
            actual_completed_at=began + timedelta(minutes=duration),
            actual_duration_minutes=duration,
            revision=5,
            updated_at=began + timedelta(minutes=duration + rng.randrange(1, 20)),
        )
        self.notify(ticket, worker_id)

    def notify(self, ticket: dict, worker_id: int):
        assigned = ticket["visit_window_start"] - timedelta(hours=self.rng.randrange(1, 14))
        assigned = max(assigned, ticket["received_at"] + timedelta(minutes=2))
        self.add(
            "notification_events",
            recipient_id=worker_id,
            ticket_id=ticket["id"],
            kind="ticket_assigned",
            data={"ticket_id": ticket["id"], "worker_id": worker_id},
            attempt_count=1,
            created_at=assigned,
            websocket_delivered_at=assigned,
            push_delivered_at=assigned + timedelta(seconds=2),
        )

    def allocate(self, ticket: dict, spec: WorkTypeSpec, brigade, appliances):
        office_id = brigade["area"]["office_id"]
        required = dict(REQUIRED_APPLIANCES.get(spec.code, ()))
        quantities = {key: quantity for key, quantity in required.items()}
        for key, low, high in appliances:
            quantity = self.rng.randrange(low, high + 1)
            if quantity:
                quantities[key] = max(quantities.get(key, 0), quantity)
        for key in sorted(quantities, key=lambda k: self.appliance_ids[k]):
            self.add(
                "ticket_appliances",
                ticket_id=ticket["id"],
                appliance_id=self.appliance_ids[key],
                office_id=office_id,
                quantity=quantities[key],
                created_at=ticket["received_at"],
            )

    def comments(self, ticket: dict, spec: WorkTypeSpec):
        rng = self.rng
        if spec.category != "emergency" and rng.random() < 0.3:
            note = rng.choice(DISPATCHER_NOTES).format(hour=rng.randrange(14, 20))
            self.add(
                "ticket_comments",
                ticket_id=ticket["id"],
                author_id=rng.randrange(1, OBSERVERS + 1),
                text=note,
                created_at=ticket["received_at"] + timedelta(minutes=rng.randrange(3, 40)),
                updated_at=ticket["received_at"] + timedelta(minutes=rng.randrange(40, 60)),
            )
        if ticket["status"] == "completed" and rng.random() < 0.6:
            report = rng.choice(WORKER_REPORTS[spec.category]).format(
                meters=rng.randrange(8, 31),
                speed=rng.choice((93, 94, 480, 910, 940)),
                wifi=rng.randrange(310, 620),
            )
            self.add(
                "ticket_comments",
                ticket_id=ticket["id"],
                author_id=ticket["assigned_worker_id"],
                text=report,
                created_at=ticket["actual_completed_at"],
                updated_at=ticket["actual_completed_at"],
            )
        if ticket["status"] == "wont_fix" and ticket["assigned_worker_id"]:
            self.add(
                "ticket_comments",
                ticket_id=ticket["id"],
                author_id=ticket["assigned_worker_id"],
                text="Клиент не открыл дверь, по телефону не ответил. Ждал 15 минут у двери.",
                created_at=ticket["updated_at"],
                updated_at=ticket["updated_at"],
            )
        elif ticket["status"] == "planned" and rng.random() < 0.05:
            brigade = next(b for b in self.brigades if b["id"] == ticket["brigade_id"])
            foreman = self.tables["brigades"][brigade["id"] - 1]["foreman_id"]
            self.add(
                "ticket_comments",
                ticket_id=ticket["id"],
                author_id=foreman,
                text=rng.choice(FOREMAN_NOTES),
                created_at=self.snapshot - timedelta(minutes=rng.randrange(5, 50)),
                updated_at=self.snapshot - timedelta(minutes=rng.randrange(0, 5)),
            )

    # -- scenarios ---------------------------------------------------------------------------

    def open_tickets(self):
        return [t for t in self.tables["tickets"] if t["status"] == "planned"]

    def shortage(self):
        """One connection of the day waits: the office has no free router for it."""
        router = self.appliance_ids[REQUIRED_APPLIANCES["connection"][0][0]]
        for ticket in self.open_tickets():
            if ticket["category"] != "connection":
                continue
            self.tables["ticket_appliances"] = [
                row
                for row in self.tables["ticket_appliances"]
                if not (row["ticket_id"] == ticket["id"] and row["appliance_id"] == router)
            ]
            ticket["description"] += (
                " Свободных роутеров SmartBox GIGA в офисе нет: остаток закрыт резервами "
                "других подключений и нормой запаса инженеров. Резерв не создан, ждём поставку."
            )
            self.package.scenarios["equipment_shortage"].append(ticket["id"])
            self.shortage_office = next(
                a["office_id"] for a in self.areas if a["id"] == ticket["service_area_id"]
            )
            return
        self.shortage_office = None

    def missing_coordinates(self):
        """A new building that the geocoder does not know yet: its points stay empty."""
        open_tickets = self.open_tickets()
        if len(open_tickets) < 8:
            return
        locations = {row["id"]: row for row in self.tables["locations"]}
        shortage = set(self.package.scenarios["equipment_shortage"])
        candidates = [t for t in open_tickets if t["id"] not in shortage]
        ticket = candidates[len(candidates) // 2]
        building_id = locations[ticket["location_id"]]["building_id"]
        for location in locations.values():
            if location["building_id"] == building_id:
                location.update(latitude=None, longitude=None)
        for row in self.tables["tickets"]:
            if locations[row["location_id"]]["building_id"] == building_id:
                self.package.scenarios["missing_coordinates"].append(row["id"])

    def overlapping(self):
        windows = defaultdict(list)
        locations = {row["id"]: row for row in self.tables["locations"]}
        for ticket in self.open_tickets():
            key = (
                locations[ticket["location_id"]]["building_id"],
                ticket["visit_window_start"],
            )
            windows[key].append(ticket["id"])
        for ids in windows.values():
            if len(ids) > 1:
                self.package.scenarios["overlapping_windows"].extend(ids)

    # -- stock -------------------------------------------------------------------------------

    def stock(self):
        reserved = defaultdict(int)
        statuses = {t["id"]: t["status"] for t in self.tables["tickets"]}
        for row in self.tables["ticket_appliances"]:
            if statuses[row["ticket_id"]] in ("planned", "in_progress"):
                reserved[(row["office_id"], row["appliance_id"])] += row["quantity"]
        kit = {self.appliance_ids[key]: quantity for key, quantity in KIT_RESERVES}
        router = self.appliance_ids[REQUIRED_APPLIANCES["connection"][0][0]]
        for area in self.areas:
            office_id = area["office_id"]
            for spec in APPLIANCES:
                appliance_id = self.appliance_ids[spec.key]
                if spec.unit == "м":
                    buffer = self.rng.randrange(300, 900)
                elif spec.type == "TOOL":
                    buffer = self.rng.randrange(2, 7)
                elif spec.key in ("rj45", "patch_utp", "patch_sc"):
                    buffer = self.rng.randrange(80, 300)
                elif not spec.is_active:
                    buffer = self.rng.randrange(3, 9)
                else:
                    buffer = self.rng.randrange(8, 30)
                if office_id == self.shortage_office and appliance_id == router:
                    buffer = 0
                stock = reserved[(office_id, appliance_id)] + kit.get(appliance_id, 0) + buffer
                self.add(
                    "appliance_stocks",
                    office_id=office_id,
                    appliance_id=appliance_id,
                    stock=stock,
                    updated_at=self.snapshot,
                )
            for appliance_id, quantity in kit.items():
                self.add(
                    "office_kit_reserves",
                    office_id=office_id,
                    appliance_id=appliance_id,
                    quantity=quantity,
                    updated_at=self.stamp,
                )

    def offices(self):
        for area in self.areas:
            office = area["area"].office
            building = self.building_id(
                city_id=1,
                area_id=area["id"],
                street=office.street,
                number=office.number,
                block=office.block,
            )
            location = self.location_id(
                building_id=building,
                entrance_id=None,
                apartment=None,
                floor=None,
                point=(office.latitude, office.longitude),
            )
            self.add(
                "offices",
                id=area["office_id"],
                location_id=location,
                name=office.name,
                service_area_id=area["id"],
            )

    def run(self) -> Package:
        self.reference()
        self.staff()
        self.demand()
        self.offices()
        self.overlapping()
        self.stock()
        for name in self.package.scenarios:
            self.package.scenarios[name] = sorted(set(self.package.scenarios[name]))
        return self.package


def generate_package(
    *,
    seed=42,
    start_date=START_DATE,
    tickets=1500,
    workers=120,
    days=7,
    area_codes=None,
    layout=DEFAULT_LAYOUT,
):
    if tickets < 8 or workers < 4 or days < 1:
        raise ValueError("Нужно не менее 8 заявок, 4 исполнителей и 1 дня")
    return Generator(
        seed=seed,
        start_date=start_date,
        tickets=tickets,
        workers=workers,
        days=days,
        area_codes=area_codes,
        layout=layout,
    ).run()


def generate_dataset(**options) -> dict:
    return generate_package(**options).tables


def package_summary(package: Package) -> dict:
    tables = package.tables
    locations = {row["id"]: row for row in tables["locations"]}
    buildings = {row["id"]: row for row in tables["buildings"]}
    areas = {row["id"]: row["name"] for row in tables["service_areas"]}
    by_area = defaultdict(lambda: defaultdict(int))
    for ticket in tables["tickets"]:
        by_area[areas[ticket["service_area_id"]]][ticket["category"]] += 1
    return {
        "tickets_by_area_and_category": {k: dict(v) for k, v in sorted(by_area.items())},
        "tickets_by_status": {
            status: sum(t["status"] == status for t in tables["tickets"])
            for status in ("planned", "completed", "wont_fix")
        },
        "addresses": {
            "buildings": len(
                {locations[t["location_id"]]["building_id"] for t in tables["tickets"]}
            ),
            "streets": len(
                {
                    buildings[locations[t["location_id"]]["building_id"]]["street_id"]
                    for t in tables["tickets"]
                }
            ),
        },
        "transport": {
            kind: sum(w["transport_type"] == kind for w in tables["workers"])
            for kind in ("car", "public_transport", "walking", "bicycle")
        },
        "scenarios": {name: len(ids) for name, ids in sorted(package.scenarios.items())},
    }


# Columns of exchange format 2, before service areas replaced districts. The Bruno
# collection loads such a package to prove that old archives still import.
LEGACY_V2_COLUMNS = {
    "appliances": (
        "name",
        "description",
        "type",
        "unit",
        "is_active",
        "created_at",
        "updated_at",
        "id",
    ),
    "cities": ("name", "id"),
    "users": ("name", "surname", "lastname", "username", "role", "created_at", "updated_at", "id"),
    "work_types": (
        "name",
        "code",
        "category",
        "default_priority",
        "travel_minutes",
        "work_minutes",
        "documents_minutes",
        "norm_minutes",
        "id",
    ),
    "worker_skills": ("skill", "id"),
    "districts": ("city_id", "name", "id"),
    "streets": ("city_id", "name", "id"),
    "work_type_planning_rules": (
        "work_type_id",
        "service_duration_source",
        "configured_by",
        "updated_at",
    ),
    "workers": ("user_id", "workshift_start", "workshift_end", "transport_type", "is_on_line"),
    "buildings": ("city_id", "street_id", "district_id", "number", "block", "id"),
    "divisions": ("district_id", "created_at", "updated_at", "id"),
    "work_type_required_appliances": ("work_type_id", "appliance_id", "quantity"),
    "work_type_required_skills": ("work_type_id", "skill_id"),
    "worker_skill_assignments": ("worker_id", "skill_id"),
    "entrances": ("building_id", "number", "id"),
    "locations": (
        "building_id",
        "entrance_id",
        "floor",
        "apartment",
        "latitude",
        "longitude",
        "id",
    ),
    "offices": ("name", "location_id", "id"),
    "tickets": (
        "location_id",
        "title",
        "description",
        "work_type",
        "work_type_id",
        "category",
        "priority",
        "received_at",
        "sla_deadline_at",
        "required_transport_type",
        "service_duration_source",
        "status",
        "lifecycle_state",
        "assigned_worker_id",
        "is_pinned",
        "visit_window_start",
        "visit_window_end",
        "planned_start_at",
        "planned_end_at",
        "estimated_duration_minutes",
        "actual_duration_minutes",
        "revision",
        "execution_cycle",
        "actual_started_at",
        "actual_completed_at",
        "cancel_reason",
        "last_event_id",
        "created_at",
        "updated_at",
        "id",
    ),
    "appliance_stocks": ("office_id", "appliance_id", "stock", "updated_at"),
    "brigades": (
        "name",
        "foreman_id",
        "office_id",
        "division_id",
        "created_at",
        "updated_at",
        "id",
    ),
    "notification_events": (
        "recipient_id",
        "ticket_id",
        "kind",
        "data",
        "created_at",
        "websocket_delivered_at",
        "push_delivered_at",
        "attempt_count",
        "next_attempt_at",
        "last_error",
        "id",
    ),
    "ticket_appliances": ("ticket_id", "appliance_id", "office_id", "quantity", "created_at"),
    "ticket_comments": ("ticket_id", "author_id", "text", "created_at", "updated_at", "id"),
    "work_events": (
        "event_type",
        "ticket_id",
        "worker_id",
        "district_id",
        "route_date",
        "occurred_at",
        "recorded_at",
        "actor_id",
        "reason",
        "previous_state",
        "new_state",
        "before_revision",
        "after_revision",
        "idempotency_key",
        "payload",
        "id",
    ),
    "worker_day_states": (
        "worker_id",
        "district_id",
        "route_date",
        "revision",
        "available",
        "unavailable_at",
        "unavailable_until",
        "last_location_id",
        "current_ticket_id",
        "current_destination_id",
        "en_route_started_at",
        "expected_available_at",
        "reason",
        "created_at",
        "updated_at",
        "id",
    ),
    "brigade_members": ("brigade_id", "worker_id"),
    "day_plan_revisions": (
        "district_id",
        "route_date",
        "revision",
        "previous_revision",
        "event_id",
        "actor_id",
        "fingerprint",
        "diff",
        "result",
        "is_current",
        "created_at",
        "id",
    ),
    "equipment_movements": (
        "ticket_id",
        "execution_cycle",
        "appliance_id",
        "office_id",
        "event_id",
        "movement",
        "quantity",
        "created_at",
        "id",
    ),
    "routes": ("worker_id", "route_date", "route_number", "geojson", "created_at", "id"),
}


def serialize_legacy_v2(tables: dict, format: str) -> bytes:
    """The same package in exchange format 2: districts instead of service areas."""
    legacy = {}
    for name, columns in LEGACY_V2_COLUMNS.items():
        legacy[name] = [
            [
                # District i of a generated package is service area 100 + i.
                encode_cell(
                    row["service_area_id"] - 100 if column == "district_id" else row.get(column)
                )
                for column in columns
            ]
            for row in tables.get(name, [])
        ]
    output = io.BytesIO()
    if format == "csv":
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            files = [("manifest.json", json.dumps({"format_version": "2"}).encode())]
            for name, rows in legacy.items():
                buffer = io.StringIO(newline="")
                writer = csv.writer(buffer, lineterminator="\r\n")
                writer.writerow(LEGACY_V2_COLUMNS[name])
                writer.writerows(rows)
                files.append((name + ".csv", buffer.getvalue().encode("utf-8-sig")))
            for name, data in files:
                info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                archive.writestr(info, data)
    else:
        workbook = Workbook(write_only=True)
        workbook.create_sheet("_meta").append(["format_version", "2"])
        for name, rows in legacy.items():
            sheet = workbook.create_sheet(name)
            sheet.append(list(LEGACY_V2_COLUMNS[name]))
            for values in rows:
                cells = []
                for value in values:
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.data_type = "s"
                    cells.append(cell)
                sheet.append(cells)
        workbook.save(output)
        workbook.close()
    return output.getvalue()


def write_dataset(output: Path, *, format_version=None, **options) -> dict:
    package = generate_package(**options)
    tables = package.tables
    output.mkdir(parents=True, exist_ok=True)
    files = {}
    counts = None
    for format, extension in (("csv", "zip"), ("xlsx", "xlsx")):
        if format_version == "2":
            content = serialize_legacy_v2(tables, format)
        else:
            content = serialize(tables, format)
        # Check every generated row with the same parser used by the upload API.
        parsed = parse_file(content, f"dataset.{extension}")
        counts = {k: len(v) for k, v in parsed.items()}
        assert all(counts[k] == len(tables[k]) for k in counts)
        path = output / f"dataset.{extension}"
        path.write_bytes(content)
        files[path.name] = {"bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()}
    manifest = {
        "generator": "generate_synthetic.py",
        "parameters": options,
        "planning_day": package.planning_day,
        **({"format_version": format_version} if format_version else {}),
        "counts": counts,
        "summary": package_summary(package),
        "scenarios": {name: ids[:50] for name, ids in sorted(package.scenarios.items())},
        "files": files,
        "notes": [
            "Адреса — реальные жилые дома Москвы из OpenStreetMap (© участники "
            "OpenStreetMap, ODbL 1.0): улица, дом, корпус, координаты подъезда, этажность "
            "и нумерация квартир. Квартира, этаж и подъезд заявки выбраны внутри "
            "реальных границ дома; сами заявки, люди и комментарии вымышлены.",
            "Последний день периода — день планирования: заявки ждут назначения. "
            "Предыдущие дни — история выполненных и отменённых визитов.",
            "Маршрутов, ревизий дня и журналов исполнения нет: их создают планировщик "
            "и приложение исполнителя.",
            "Импортированные учётные записи не имеют известного пароля; пароль задаётся отдельно.",
        ],
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=json_default) + "\n",
        encoding="utf-8",
    )
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("../data/synthetic/standard"))
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--date", type=date.fromisoformat, default=START_DATE)
    parser.add_argument("--tickets", type=int, default=1500)
    parser.add_argument("--workers", type=int, default=120)
    parser.add_argument("--days", type=int, default=7)
    parser.add_argument(
        "--format-version", choices=("2",), help="Старый формат обмена (для проверки импорта)"
    )
    args = parser.parse_args()
    result = write_dataset(
        args.output,
        format_version=args.format_version,
        seed=args.seed,
        start_date=args.date,
        tickets=args.tickets,
        workers=args.workers,
        days=args.days,
    )
    print(json.dumps(result["counts"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
